import os
import re
import subprocess
import unicodedata
import warnings

import numpy as np
import pandas as pd
from openpyxl import load_workbook


EXCEL_HEADER_SCAN_ROWS = 15
EXCELCNV_CANDIDATE_PATHS = (
    r"C:\Program Files\Microsoft Office\root\Office16\excelcnv.exe",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\excelcnv.exe",
)

KEYWORD_RULES = {
    "Tên hoạt chất": ["hoạt chất"],
    "Tên thuốc": ["tên thuốc"],
    "Nồng độ, hàm lượng": ["nồng độ", "hàm lượng"],
    "Số lượng": ["số lượng"],
    "Số đăng ký": ["số đăng ký"],
    "GĐKLH hoặc GPNK": ["gđklh", "gpnk"],
    "Đơn giá trúng thầu (VND)": ["đơn giá"],
    "Thành tiền (VND)": ["thành tiền"],
    "Nhà thầu trúng thầu": ["nhà thầu"],
    "Danh mục hàng hóa": ["tên hàng", "hàng hóa", "danh mục hàng"],
    "Tên phần/lô": ["tên phần", "tên lô"],
    "Mặt hàng dự thầu": ["mặt hàng dự thầu", "mặt hàng"],
    "Ký mã hiệu": ["ký mã", "mã hiệu"],
    "Tính năng kỹ thuật": ["tính năng", "kỹ thuật"],
    "Xuất xứ": ["xuất xứ", "nước sản xuất"],
    "Hãng sản xuất": ["hãng sản xuất"],
    "Năm sản xuất": ["năm sản xuất"],
}

warnings.filterwarnings(
    "ignore",
    message=r"Cannot parse header or footer so it will be ignored",
    category=UserWarning,
    module="openpyxl.worksheet.header_footer",
)
warnings.filterwarnings(
    "ignore",
    message=r"Unknown extension is not supported and will be removed",
    category=UserWarning,
    module=r"openpyxl\.worksheet\._reader",
)
warnings.filterwarnings(
    "ignore",
    message=r"Conditional Formatting extension is not supported and will be removed",
    category=UserWarning,
    module=r"openpyxl\.worksheet\._reader",
)


def clean_col_str(s):
    if not isinstance(s, str):
        s = str(s)
    s = unicodedata.normalize("NFC", s)
    return re.sub(r"\s+", " ", s).strip().lower()


def normalize_header_lookup_key(value):
    text = clean_col_str(value)
    if not text:
        return ""
    text = text.replace("vnđ", "vnd")
    text = text.replace("đồng", "dong")
    text = re.sub(r"[\W_]+", "", text, flags=re.UNICODE)
    return text


def build_schema_mapping_config(config):
    mapping = {}
    for canonical_col in config.get("db_mapping", {}).keys():
        mapping[canonical_col] = canonical_col
    mapping.update(config.get("column_mapping", {}))
    return mapping


def _is_unit_price_target(value) -> bool:
    return clean_col_str(value) == "đơn giá trúng thầu (vnd)"


def _is_unit_price_source(value) -> bool:
    return "đơn giá" in clean_col_str(value)


def _is_ambiguous_price_source(value) -> bool:
    source_clean = clean_col_str(value)
    return "giá" in source_clean and "đơn giá" not in source_clean


def get_smart_column_mapping(df_columns, mapping_config):
    final_map = {}
    clean_mapping_config = {clean_col_str(k): v for k, v in mapping_config.items()}
    normalized_mapping_config = {
        normalize_header_lookup_key(k): v
        for k, v in mapping_config.items()
        if normalize_header_lookup_key(k)
    }
    best_target_choice = {}

    def resolve_explicit_mapping(col):
        col_clean = clean_col_str(col)
        if col in mapping_config:
            return mapping_config[col]
        if col_clean in clean_mapping_config:
            return clean_mapping_config[col_clean]
        col_lookup = normalize_header_lookup_key(col)
        if col_lookup and col_lookup in normalized_mapping_config:
            return normalized_mapping_config[col_lookup]
        return None

    unit_price_locked_by_explicit_column = any(
        _is_unit_price_target(resolve_explicit_mapping(col))
        and _is_unit_price_source(col)
        for col in df_columns
    )

    def resolve_contextual_target(source_col, target_col):
        if (
            unit_price_locked_by_explicit_column
            and _is_unit_price_target(target_col)
            and _is_ambiguous_price_source(source_col)
        ):
            return None
        return target_col

    def register_candidate(source_col, target_col, priority):
        target_col = resolve_contextual_target(source_col, target_col)
        if not target_col:
            return
        source_lookup = normalize_header_lookup_key(source_col)
        target_lookup = normalize_header_lookup_key(target_col)
        canonical_exact = int(bool(source_lookup) and source_lookup == target_lookup)
        unit_price_signal = int(_is_unit_price_target(target_col) and _is_unit_price_source(source_col))
        ambiguous_price_signal = int(
            _is_unit_price_target(target_col)
            and _is_ambiguous_price_source(source_col)
        )
        candidate = (
            priority,
            canonical_exact,
            unit_price_signal,
            -ambiguous_price_signal,
            len(str(source_col or "")),
        )
        current = best_target_choice.get(target_col)
        if current is None or candidate > current[0]:
            best_target_choice[target_col] = (candidate, source_col)

    for col in df_columns:
        col_clean = clean_col_str(col)
        if col in mapping_config:
            register_candidate(col, mapping_config[col], 3)
            continue
        if col_clean in clean_mapping_config:
            register_candidate(col, clean_mapping_config[col_clean], 3)
            continue
        col_lookup = normalize_header_lookup_key(col)
        if col_lookup and col_lookup in normalized_mapping_config:
            register_candidate(col, normalized_mapping_config[col_lookup], 3)
            continue

        for target_col, keywords in KEYWORD_RULES.items():
            if any(kw in col_clean for kw in keywords):
                register_candidate(col, target_col, 1)
                break

    for target_col, (_, source_col) in best_target_choice.items():
        final_map[source_col] = target_col
    return final_map


def collapse_duplicate_columns(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty or not df.columns.duplicated().any():
        return df

    collapsed = pd.DataFrame(index=df.index)
    for col_name in df.columns.unique():
        same_name = df.loc[:, df.columns == col_name]
        if isinstance(same_name, pd.Series):
            collapsed[col_name] = same_name
            continue

        merged = same_name.iloc[:, 0]
        for idx in range(1, same_name.shape[1]):
            merged = merged.combine_first(same_name.iloc[:, idx])
        collapsed[col_name] = merged

    return collapsed


def _normalize_numeric_text(value) -> str:
    if pd.isna(value):
        return ""

    text = str(value).strip()
    if text in {"nan", "None", "<NA>", "NaT", "nat", "null", "NULL"}:
        return ""

    text = text.replace("\u00a0", " ")
    text = re.sub(r"[^\d,.\-]", "", text)
    if not text or text in {"-", ".", ",", "-.", "-,"}:
        return ""

    sign = ""
    if text.startswith("-"):
        sign = "-"
        text = text[1:]
    text = text.replace("-", "")

    comma_count = text.count(",")
    dot_count = text.count(".")

    if comma_count and dot_count:
        decimal_sep = "," if text.rfind(",") > text.rfind(".") else "."
        thousand_sep = "." if decimal_sep == "," else ","
        text = text.replace(thousand_sep, "")
        if decimal_sep == ",":
            text = text.replace(",", ".", 1).replace(",", "")
    elif comma_count:
        parts = text.split(",")
        if comma_count > 1 and all(len(part) == 3 for part in parts[1:]):
            text = "".join(parts)
        elif len(parts[-1]) == 3 and len(parts[0]) <= 3:
            text = "".join(parts)
        else:
            text = text.replace(",", ".", 1).replace(",", "")
    elif dot_count:
        parts = text.split(".")
        if dot_count > 1 and all(len(part) == 3 for part in parts[1:]):
            text = "".join(parts)
        elif len(parts[-1]) == 3 and len(parts[0]) <= 3:
            text = "".join(parts)
        else:
            text = text.replace(".", ".", 1).replace(".", "", dot_count - 1)

    return sign + text


def clean_numeric_series(series: pd.Series) -> pd.Series:
    s = series.map(_normalize_numeric_text)
    return pd.to_numeric(s, errors="coerce")


def drop_header_legend_rows(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return df

    def is_parenthesized_ordinal(value) -> bool:
        if pd.isna(value):
            return False
        text = str(value).strip()
        return bool(re.fullmatch(r"\(\d+\)", text))

    def parse_numeric_ordinal(value):
        if pd.isna(value):
            return None
        text = str(value).strip()
        if not text:
            return None
        if re.fullmatch(r"\d+(?:\.0+)?", text):
            return int(float(text))
        return None

    def is_sequential_numeric_legend(values) -> bool:
        numeric_values = [parse_numeric_ordinal(value) for value in values]
        if any(value is None for value in numeric_values):
            return False
        if len(numeric_values) < 3:
            return False
        if len(set(numeric_values)) != len(numeric_values):
            return False
        return all(
            current == previous + 1
            for previous, current in zip(numeric_values, numeric_values[1:])
        )

    mask = []
    for _, row in df.iterrows():
        values = [value for value in row.tolist() if not pd.isna(value) and str(value).strip() != ""]
        if len(values) >= 3 and (
            all(is_parenthesized_ordinal(value) for value in values)
            or is_sequential_numeric_legend(values)
        ):
            mask.append(False)
        else:
            mask.append(True)

    return df.loc[mask].copy()


def drop_invalid_value_rows(df: pd.DataFrame, schema_name: str) -> pd.DataFrame:
    if df is None or df.empty:
        return df

    price_col = "Đơn giá trúng thầu (VND)"
    amount_col = "Thành tiền (VND)"
    quantity_col = "Số lượng" if schema_name == "MEDICINE_STANDARD" else "Khối lượng"

    if price_col not in df.columns:
        return df

    df = df.copy()

    if quantity_col in df.columns:
        df[quantity_col] = clean_numeric_series(df[quantity_col])
    if price_col in df.columns:
        df[price_col] = clean_numeric_series(df[price_col])
    if amount_col in df.columns:
        df[amount_col] = clean_numeric_series(df[amount_col])
    elif quantity_col in df.columns and price_col in df.columns:
        df[amount_col] = np.nan

    if all(col in df.columns for col in [quantity_col, price_col, amount_col]):
        mask_missing = df[amount_col].isna()
        mask_has_inputs = (
            df[quantity_col].notna()
            & (df[quantity_col] != 0)
            & df[price_col].notna()
            & (df[price_col] != 0)
        )
        df.loc[mask_missing & mask_has_inputs, amount_col] = (
            df.loc[mask_missing & mask_has_inputs, quantity_col]
            * df.loc[mask_missing & mask_has_inputs, price_col]
        )

    quantity_series = (
        df[quantity_col]
        if quantity_col in df.columns
        else pd.Series([np.nan] * len(df), index=df.index)
    )
    price_series = (
        df[price_col]
        if price_col in df.columns
        else pd.Series([np.nan] * len(df), index=df.index)
    )
    invalid_mask = (
        quantity_series.isna()
        | (quantity_series <= 0)
        | price_series.isna()
        | (price_series <= 0)
    )

    return df.loc[~invalid_mask].copy()


def detect_excel_header_index(file_path, scan_rows=EXCEL_HEADER_SCAN_ROWS):
    readable_path = resolve_excel_readable_path(file_path)
    best_sheet = detect_best_excel_sheet(readable_path, scan_rows=scan_rows)
    df_temp = pd.read_excel(readable_path, sheet_name=best_sheet, header=None, nrows=scan_rows)
    if df_temp.empty:
        raise ValueError("File Excel rỗng")

    non_blank_counts = df_temp.notna().sum(axis=1)
    if non_blank_counts.empty or int(non_blank_counts.max()) <= 0:
        raise ValueError("Không xác định được dòng header")

    return int(non_blank_counts.idxmax())


def load_excel_with_detected_header(file_path, sample_rows=None, dtype=None, scan_rows=EXCEL_HEADER_SCAN_ROWS):
    readable_path = resolve_excel_readable_path(file_path)
    best_sheet = detect_best_excel_sheet(readable_path, scan_rows=scan_rows)
    header_idx = detect_excel_header_index(readable_path, scan_rows=scan_rows)
    read_kwargs = {"header": header_idx, "sheet_name": best_sheet}
    if sample_rows is not None:
        read_kwargs["nrows"] = sample_rows
    if dtype is not None:
        read_kwargs["dtype"] = dtype
        # Preserve literal placeholders like "NA" from Excel instead of
        # silently coercing them to missing values.
        read_kwargs["keep_default_na"] = False
    return pd.read_excel(readable_path, **read_kwargs)


def count_excel_rows_with_detected_header(file_path, scan_rows=EXCEL_HEADER_SCAN_ROWS):
    df = load_excel_with_detected_header(file_path, scan_rows=scan_rows)
    return len(df.dropna(how="all"))


def get_excel_sheet_name_groups(file_path):
    readable_path = resolve_excel_readable_path(file_path)

    try:
        workbook = load_workbook(readable_path, read_only=True, data_only=True)
    except Exception:
        with pd.ExcelFile(readable_path) as workbook:
            sheet_names = list(workbook.sheet_names or [])
        return {
            "all": sheet_names,
            "visible": sheet_names,
            "hidden": [],
        }

    try:
        all_sheets = []
        visible_sheets = []
        hidden_sheets = []

        for worksheet in workbook.worksheets:
            sheet_name = worksheet.title
            sheet_state = getattr(worksheet, "sheet_state", "visible")
            all_sheets.append(sheet_name)
            if sheet_state == "visible":
                visible_sheets.append(sheet_name)
            else:
                hidden_sheets.append(sheet_name)

        return {
            "all": all_sheets,
            "visible": visible_sheets,
            "hidden": hidden_sheets,
        }
    finally:
        workbook.close()


def detect_best_excel_sheet(file_path, scan_rows=EXCEL_HEADER_SCAN_ROWS):
    readable_path = resolve_excel_readable_path(file_path)
    with pd.ExcelFile(readable_path) as workbook:
        best_sheet = None
        best_score = (-1, -1)
        sheet_groups = get_excel_sheet_name_groups(readable_path)
        candidate_sheets = sheet_groups["visible"] or sheet_groups["all"] or list(workbook.sheet_names or [])

        for sheet_name in candidate_sheets:
            try:
                preview_df = pd.read_excel(workbook, sheet_name=sheet_name, header=None, nrows=scan_rows)
            except Exception:
                continue
            if preview_df.empty:
                continue

            non_blank_per_row = preview_df.notna().sum(axis=1)
            max_non_blank = int(non_blank_per_row.max()) if not non_blank_per_row.empty else 0
            total_non_blank = int(preview_df.notna().sum().sum())
            score = (max_non_blank, total_non_blank)

            if score > best_score:
                best_score = score
                best_sheet = sheet_name

        if best_sheet is None:
            raise ValueError("Không tìm thấy sheet Excel có dữ liệu")
        return best_sheet


def convert_legacy_xls_to_xlsx(xls_path):
    xls_path = os.path.abspath(str(xls_path))
    xlsx_path = os.path.splitext(xls_path)[0] + ".xlsx"
    if os.path.exists(xlsx_path) and os.path.getmtime(xlsx_path) >= os.path.getmtime(xls_path):
        return xlsx_path

    excelcnv_path = next((path for path in EXCELCNV_CANDIDATE_PATHS if os.path.exists(path)), None)
    if excelcnv_path:
        output_dir = os.path.dirname(xls_path)
        command = [excelcnv_path, "-oice", xls_path, "-nme", output_dir]
        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                check=False,
                timeout=120,
            )
        except Exception:
            result = None
        if result and result.returncode == 0 and os.path.exists(xlsx_path):
            return xlsx_path

    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        if os.path.exists(xlsx_path):
            return xlsx_path
        raise ImportError(
            "Missing optional dependency 'xlrd', và cũng không có excelcnv.exe hoặc win32com/pywin32 để convert file .xls."
        ) from exc

    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(xls_path, ReadOnly=True)
        workbook.SaveAs(xlsx_path, FileFormat=51)
        return xlsx_path if os.path.exists(xlsx_path) else None
    except Exception:
        if os.path.exists(xlsx_path):
            return xlsx_path
        raise
    finally:
        try:
            if workbook is not None:
                workbook.Close(False)
        except Exception:
            pass
        try:
            if excel is not None:
                excel.Quit()
        except Exception:
            pass


def resolve_excel_readable_path(file_path):
    normalized_path = os.path.abspath(str(file_path))
    if os.path.splitext(normalized_path)[1].lower() != ".xls":
        return normalized_path
    converted_path = convert_legacy_xls_to_xlsx(normalized_path)
    return converted_path or normalized_path
